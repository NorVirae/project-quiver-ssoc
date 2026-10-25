"""Build the Word handbook from the same Markdown used by MkDocs."""
from pathlib import Path
import argparse
import json
import re
import markdown
from lxml import html
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
ORDER = json.loads((ROOT / 'course-order.json').read_text(encoding='utf-8'))

def el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items(): e.set(qn('w:' + k), str(v))
    return e

def bookmark(p, name, number):
    p._p.insert(0, el('w:bookmarkStart', id=number, name=name))
    p._p.append(el('w:bookmarkEnd', id=number))

def field(p, instruction):
    r = p.add_run()
    r._r.append(el('w:fldChar', fldCharType='begin'))
    t = el('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = ' ' + instruction + ' '
    r._r.append(t)
    r._r.append(el('w:fldChar', fldCharType='end'))

def text_run(p, text, bold=False, italic=False, code=False):
    if not text: return
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if code: r.font.name = 'Consolas'; r.font.size = Pt(9)
    return r

def hyperlink(p, label, target):
    h = OxmlElement('w:hyperlink')
    if target.endswith('.md') or '.md#' in target:
        h.set(qn('w:anchor'), 'p_' + Path(target.split('#')[0]).stem.replace('-', '_'))
    elif target.startswith('http'):
        h.set(qn('r:id'), p.part.relate_to(target, RT.HYPERLINK, is_external=True))
    else:
        return text_run(p, label)
    r = el('w:r'); props = el('w:rPr')
    props.append(el('w:color', val='126B78')); props.append(el('w:u', val='single'))
    r.append(props); t = el('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)

def inline(p, node, bold=False, italic=False):
    text_run(p, node.text, bold, italic)
    for child in node:
        if child.tag == 'a': hyperlink(p, ''.join(child.itertext()), child.get('href', ''))
        elif child.tag == 'strong': inline(p, child, True, italic)
        elif child.tag == 'em': inline(p, child, bold, True)
        elif child.tag == 'code': text_run(p, child.text, bold, italic, True)
        elif child.tag == 'br': p.add_run().add_break()
        elif child.tag != 'img': inline(p, child, bold, italic)
        text_run(p, child.tail, bold, italic)

def set_cell(cell):
    pr = cell._tc.get_or_add_tcPr()
    margins = el('w:tcMar')
    for side in ['top', 'bottom']: margins.append(el('w:' + side, w=65, type='dxa'))
    for side in ['left', 'right']: margins.append(el('w:' + side, w=90, type='dxa'))
    pr.append(margins)

def render_table(doc, node):
    rows = node.xpath('.//tr')
    if not rows: return
    n = len(rows[0]); table = doc.add_table(rows=0, cols=n); table.autofit = False
    first_header = ''.join(rows[0][0].itertext())
    three_column_widths = {'Team': [1.45,3.2,2.45], 'Checkpoint': [1.65,3.6,1.85]}
    widths = ([2.02, 5.08] if n == 2 else three_column_widths.get(first_header, [0.92,5.33,0.85]) if n == 3 else [7.1/n]*n)
    for c, width in zip(table.columns, widths): c.width = Inches(width)
    borders = el('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        borders.append(el('w:' + side, val='single', sz=4, color='D4DADD'))
    table._tbl.tblPr.append(borders)
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for c, width, source in zip(cells, widths, row):
            c.width = Inches(width); set_cell(c)
            p = c.paragraphs[0]; p.style = doc.styles['Table Text']; inline(p, source, idx == 0)
            if idx == 0: c._tc.get_or_add_tcPr().append(el('w:shd', fill='F0F2F3'))
        trpr = table.rows[-1]._tr.get_or_add_trPr(); trpr.append(el('w:cantSplit'))
        if idx == 0: trpr.append(el('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def picture(doc, src, alt, cover=False):
    path = ROOT / 'docs' / src
    with Image.open(path) as im: ratio = im.height / im.width
    width = min(7.1, (3.25 if cover else 1.72) / ratio)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(5)
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set('descr', alt)

def render_page(doc, item, number):
    body = (ROOT / 'docs' / (item['slug'] + '.md')).read_text(encoding='utf-8')
    tree = html.fragment_fromstring(markdown.markdown(body, extensions=['tables']), create_parent='div')
    for node in tree:
        tag = node.tag
        if tag == 'h1':
            if item['slug'] == 'index':
                p = doc.add_paragraph('University of Lagos student course', 'Subtitle')
                p = doc.add_paragraph(node.text, 'Title')
            else:
                p = doc.add_paragraph(item['group'].upper(), 'Eyebrow')
                p.paragraph_format.page_break_before = True
                p = doc.add_paragraph(node.text, 'Heading 1')
            bookmark(p, 'p_' + item['slug'].replace('-', '_'), number)
        elif tag in ['h2', 'h3']:
            doc.add_paragraph(''.join(node.itertext()), 'Heading 2')
        elif tag == 'p':
            imgs = node.findall('img')
            if imgs:
                for img in imgs: picture(doc, img.get('src'), img.get('alt',''), item['slug'] == 'index')
            else:
                p = doc.add_paragraph(); inline(p, node)
                if node.find('em') is not None and len(''.join(node.itertext())) < 190:
                    p.style = doc.styles['Caption']
        elif tag in ['ul', 'ol']:
            for i, li in enumerate(node.findall('li'), 1):
                p = doc.add_paragraph(style='List Item')
                raw = ''.join(li.itertext())
                if raw.startswith('[ ] '):
                    text_run(p, '☐  ')
                    if li.text: li.text = li.text[4:]
                else: text_run(p, (str(i) + '.  ') if tag == 'ol' else '•  ')
                inline(p, li)
        elif tag == 'table': render_table(doc, node)
        elif tag == 'pre':
            p = doc.add_paragraph(style='Code'); text_run(p, ''.join(node.itertext()), code=True)
        elif tag == 'hr':
            doc.add_paragraph('_' * 66)

def build(output):
    doc = Document(); sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.66), Inches(.64)
    sec.left_margin = sec.right_margin = Inches(.7)
    sec.header_distance = sec.footer_distance = Inches(.3)
    st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11)
    st.font.color.rgb = RGBColor(0,0,0); st.paragraph_format.line_spacing = Pt(14)
    st.paragraph_format.space_after = Pt(5); st.paragraph_format.widow_control = True
    for name, size in [('Title',27),('Heading 1',19),('Heading 2',12.5),('Subtitle',12)]:
        s = doc.styles[name]; s.font.name='Arial'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.space_before=Pt(9 if name=='Heading 2' else 0)
        s.paragraph_format.space_after=Pt(5 if name=='Heading 2' else 9)
        s.paragraph_format.line_spacing=1.05
        s.font.bold = name.startswith('Heading')
    for name, size in [('Eyebrow',9),('List Item',11),('Table Text',10.5),('Code',9)]:
        s = doc.styles.add_style(name, 1); s.base_style=doc.styles['Normal']; s.font.size=Pt(size)
    doc.styles['Eyebrow'].font.color.rgb=RGBColor.from_string('536069')
    doc.styles['Eyebrow'].paragraph_format.keep_with_next=True
    s=doc.styles['List Item']; s.paragraph_format.left_indent=Inches(.18); s.paragraph_format.first_line_indent=Inches(-.18); s.paragraph_format.space_after=Pt(4)
    s=doc.styles['Table Text']; s.paragraph_format.line_spacing=Pt(12.5); s.paragraph_format.space_after=Pt(0)
    s=doc.styles['Caption']; s.font.name='Arial'; s.font.size=Pt(9); s.font.color.rgb=RGBColor.from_string('46545C'); s.paragraph_format.line_spacing=Pt(11)
    s.font.bold=False; s.font.italic=False
    for s in doc.styles:
        for border in s.element.findall('.//' + qn('w:pBdr')):
            border.getparent().remove(border)
    sec.header.paragraphs[0].text='QUIVER   /   STUDENT LABORATORY HANDBOOK'
    sec.header.paragraphs[0].style='Eyebrow'
    p=sec.footer.paragraphs[0]; p.paragraph_format.tab_stops.add_tab_stop(Inches(7.1), WD_TAB_ALIGNMENT.RIGHT)
    r=p.add_run('Version 5.1  •  25 October 2026\t'); r.font.size=Pt(9)
    field(p,'PAGE'); p.style=doc.styles['Caption']
    doc.core_properties.title='Quiver student laboratory handbook'
    doc.core_properties.subject='Daily laboratory practice, equipment induction and supervised build course'
    doc.core_properties.author='Quiver course development'
    doc.core_properties.version='5.1'
    settings=doc.settings.element; settings.append(el('w:updateFields', val='true'))
    # A shorter front page introduces the handbook; complete orientation remains on the website.
    doc.add_paragraph('University of Lagos student course', 'Subtitle')
    p=doc.add_paragraph('Quiver student\nlaboratory handbook','Title'); bookmark(p,'p_index',1)
    doc.add_paragraph('Daily practice  •  Equipment induction  •  Parallel build work','Subtitle')
    picture(doc,'assets/quiver.jpg','Quiver reference aircraft from the project repository',True)
    doc.add_paragraph('Reference aircraft from the project repository. Remove all propellers for classroom and bench work.','Caption')
    doc.add_heading('Learn the routine then build',2)
    doc.add_paragraph('This course teaches students to work safely, care for shared equipment and complete a supervised Quiver build. Start with daily laboratory practice, complete the common induction, then join one of six build teams.')
    doc.add_paragraph('For each activity, read the procedure, watch the instructor demonstrate, practise under supervision and show the required evidence. Use the opening and closing checklist at every session.')
    doc.add_heading('Course outcomes',2)
    doc.add_paragraph('Identify hazards and isolate equipment; manage tools and components; use authorised lab equipment; explain Quiver systems; contribute a checked subsystem; run and evaluate a simulation; and hand over traceable work.')
    doc.add_paragraph('Local adoption is required before practical teaching. The department completes the room arrangements and approves equipment procedures. Course completion, equipment authorisation and real flight approval are separate decisions.')
    doc.add_page_break()
    doc.add_heading('Find a procedure',1)
    doc.add_paragraph('Follow the course in order during induction. Later, use this guide or the website search to find the activity you are about to perform.')
    current=None
    groups={'daily-rules':'Daily lab practice','systems':'Build and simulate','assessment':'Checklists and records','delivery':'Teaching and sources'}
    for item in ORDER[1:]:
        if item['slug'] in groups:
            doc.add_heading(groups[item['slug']],2)
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=Pt(13)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.05),WD_TAB_ALIGNMENT.RIGHT)
        hyperlink(p,item['title'],item['slug']+'.md'); p.add_run('\t'); field(p,'PAGEREF p_'+item['slug'].replace('-','_'))
    for number,item in enumerate(ORDER[1:],2):
        render_page(doc,item,number)
    output.parent.mkdir(parents=True,exist_ok=True); doc.save(output)
    print(str(output))

if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--output',type=Path,default=ROOT/'docs'/'downloads'/'Quiver_Student_Laboratory_Handbook_v5_1.docx')
    build(parser.parse_args().output)
